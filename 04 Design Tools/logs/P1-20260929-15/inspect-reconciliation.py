from openpyxl import load_workbook

previous = load_workbook('04 Design Tools/logs/P1-20260929-15/test-sheet-after.xlsx', read_only=True, data_only=False)
book = load_workbook('04 Design Tools/logs/P1-20260929-15/test-sheet-reconciliation.xlsx', read_only=True, data_only=False)
def stable(value):
    if value.__class__.__name__ == 'ArrayFormula':
        return ('ArrayFormula', value.text, value.ref)
    return value

assert previous.sheetnames == book.sheetnames
for name in book.sheetnames:
    old_rows = [[stable(value) for value in row] for row in previous[name].iter_rows(values_only=True)]
    new_rows = [[stable(value) for value in row] for row in book[name].iter_rows(values_only=True)]
    assert old_rows == new_rows, name
    print(name, 'exported formulas/values unchanged', len(new_rows), 'rows')
for name, rows in [('ADD REQUESTS', [1, 63, 64, 65, 66, 67, 68, 69]), ('GAME GUIDE BOOKS', [1, 20, 21, 22])]:
    sheet = book[name]
    print(name)
    for row in rows:
        print(row, [sheet.cell(row, col).value for col in range(1, 18)])
