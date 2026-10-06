import openpyxl,json,sys
from pathlib import Path
p=Path(__file__).parent
final=sys.argv[1]
f=openpyxl.load_workbook(p/final,data_only=False);v=openpyxl.load_workbook(p/final,data_only=True)
b=openpyxl.load_workbook(p/'typed-sales-before.xlsx',data_only=False)
def normalize(x):
 if hasattr(x,'text'):return {'formula':x.text,'ref':x.ref}
 if hasattr(x,'isoformat'):return x.isoformat()
 return x
for row in b['SALES'].iter_rows(min_row=1,max_row=53,max_col=9):
 for c in row:assert normalize(c.value)==normalize(f['SALES'][c.coordinate].value),(c.coordinate,'protected baseline changed')
tables=[t for t in f['SALES'].tables.values()]
assert len(tables)==1 and tables[0].ref=='A1:I53' and tables[0].totalsRowCount==1
raw_orders=[r for r in list(v['ORDERS'].values)[1:] if r[0]]
orders={r[0]:r for r in raw_orders}
assert len(orders)==len(raw_orders),'Duplicate Order ID'
lines=[r for r in list(v['ORDER LINES'].values)[1:] if r[0]]
sales=[(row[0].row,[c.value for c in row]) for row in v['SALES'] if row[0].value and str(row[0].value).startswith('OWA-')]
assert len(set(r[9] for _,r in sales))==len(sales),'Duplicate SALES Line ID'
assert all(n>1000 for n,r in sales),'Inside baseline table/spill/capacity'
assert orders['OWA-20261004-01'][1]=='SOLD' and orders['OWA-20261004-02'][1]=='CANCELLED'
assert orders['OWA-20261004-03'][1]=='SOLD'
assert not any(r[0]=='OWA-20261004-02' for _,r in sales),'Cancelled sale exists'
for _,r in sales:
 assert r[6]==r[4]-r[3]-r[5],('Profit mismatch',r)
 assert isinstance(r[6],(float,int))
events=[r for r in list(v['ORDER REQUESTS'].values)[1:] if r[0]]
for rid in ['w1-20261004-recovery-create','w1-20261004-recovery-confirm']:
 states=[r[7] for r in events if r[2]==rid]
 assert 'NEEDS_REVIEW' in states and states[-1]=='DONE'
cancel=[r for r in events if r[2]=='w1-1791059465888-wzpaj3c0gl']
assert len([r for r in cancel if r[7]=='PREPARED'])==1 and len([r for r in cancel if r[7]=='DONE'])==1
assert len(set(r[0] for r in events))==len(events)
assert all(r[1] in ['RESERVED','SOLD','RELEASED'] for r in [(x[1],x[12]) for x in lines])
copies=[(r[0].value,r[2].value,r[26].value) for r in v['GAME GUIDE BOOKS'] if str(r[0].value).startswith('W1 QA')]
assert len(copies)==8 and len([x for x in copies if x[2]])==len(set(x[2] for x in copies if x[2]))
out={'result':'PASS','export':final,'orders':{k:{'status':r[1],'channel':r[2],'customerShipping':r[8],'shippingSubsidy':r[9],'requestId':r[13]} for k,r in orders.items()},'sales':[{'row':n,'order':r[0],'price':r[4],'cost':r[3],'subsidy':r[5],'profit':r[6],'lineId':r[9]} for n,r in sales],'copies':copies,'table':tables[0].ref,'finiteG2AndLegacy52':'unchanged','cancelSameId':'one PREPARED + one DONE','injectedHoldSales':'NEEDS_REVIEW then DONE same IDs'}
(p/'google-assertions.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(out,ensure_ascii=False))
