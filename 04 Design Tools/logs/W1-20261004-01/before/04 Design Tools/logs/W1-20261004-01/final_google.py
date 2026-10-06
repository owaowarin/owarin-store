import openpyxl, json, csv
from pathlib import Path
p=Path(__file__).parent
f=openpyxl.load_workbook(p/'google-final.xlsx',data_only=True)
orders={r[0]:r for r in list(f['ORDERS'].values)[1:] if r[0]}
assert len(orders)==6
assert [orders[f'OWA-20261004-{i:02}'][1] for i in range(1,7)]==['SOLD','CANCELLED','SOLD','SOLD','SOLD','CANCELLED']
sales=[r for r in f['SALES'].values if r[0] and str(r[0]).startswith('OWA-')]
assert [(r[0],r[4],r[5],r[6]) for r in sales]==[('OWA-20261004-01',390,10,280),('OWA-20261004-03',390,0,290),('OWA-20261004-04',391,10,281),('OWA-20261004-05',400,20,280)]
lines=[r for r in list(f['ORDER LINES'].values)[1:] if r[0]]
assert len(lines)==6 and len({r[1] for r in lines})==6
assert all(r[12]==('RELEASED' if orders[r[0]][1]=='CANCELLED' else 'SOLD') for r in lines)
events=[r for r in list(f['ORDER REQUESTS'].values)[1:] if r[0]]
latest={r[2]:r[7] for r in events}
assert all(s=='DONE' for s in latest.values()),latest
assert not any(r[2]=='w1-1791060373267-p9hui5vkdu' for r in events)
assert f['GAME GUIDE BOOKS']['M31'].value==390
copies=[r for r in f['GAME GUIDE BOOKS'].values if str(r[0]).startswith('W1 QA')]
assert [r[2] for r in copies]==['Instock','Instock','Instock','Sold','Sold','Sold','Sold','Auction']
assert all(r[26] for r in copies if r[2]=='Sold')
baseline=openpyxl.load_workbook(p/'test-before.xlsx',data_only=False)
current=openpyxl.load_workbook(p/'google-final.xlsx',data_only=False)
def norm(x):
    if hasattr(x,'text'):return (x.text,x.ref)
    return x
for tab in ['GAME GUIDE BOOKS','MAGAZINE','ADD REQUESTS']:
    for row in baseline[tab]:
        for c in row:
            assert norm(c.value)==norm(current[tab][c.coordinate].value),(tab,c.coordinate)
with (p/'changes.csv').open(encoding='utf-8',newline='') as stream:
    records=list(csv.DictReader(stream))
assert all(r['change_id']=='W1-20261004-01' and all(r[k] for k in ['request_id','action','result','recovery']) for r in records)
out={'result':'PASS','source':'existing final export; no new Google writes or redundant runtime rerun','sixOrderStatuses':True,'fourExactSalesAndProfits':True,'cancel02And06NoSales':True,'claimsOwnedAndReleased':True,'latestRequestStates':latest,'externalConfirmPreflightNoJournal':True,'M31Restored390':True,'eightCopyStatuses':True,'initialP1InventoryMagazineAddJournalPreserved':True,'csvEntries':len(records)}
(p/'final-google-checks.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print('PASS final six orders / four exact ledgers / all latest requests DONE / original P1 cells preserved')
