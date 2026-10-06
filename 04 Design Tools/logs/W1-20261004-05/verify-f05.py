import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

root = Path(__file__).parent / 'attempts' / 'F05'
before = load_workbook(root / 'session42-before.xlsx', data_only=False)
after = load_workbook(root / 'session42-repaired.xlsx', data_only=False)
spec = json.loads((root / 'gap-repair-intent.json').read_text(encoding='utf8'))
def rows(book, tab):
    return [[('ARRAY_FORMULA', x.text, x.ref) if isinstance(x, ArrayFormula) else x if x is not None else '' for x in row] for row in book[tab].values]
assert before.sheetnames == after.sheetnames
for tab in before.sheetnames:
    if tab == 'ORDER REQUESTS':
        continue
    assert rows(before, tab) == rows(after, tab), ('changed-other-tab', tab)
old = rows(before, 'ORDER REQUESTS')
new = rows(after, 'ORDER REQUESTS')
gaps = {g['row']: g['eventId'] for g in spec['gaps']}
assert len(gaps) == 49 and len(set(gaps.values())) == 49
for physical in range(2, 1122):
    prior, current = old[physical - 1], new[physical - 1]
    if physical in gaps:
        assert all(v == '' for v in prior), ('prior-not-blank', physical)
        assert current[0] == gaps[physical] and current[2] == f'qa-gap-20261004-R{physical}'
        assert current[4] == 'QA_CAPACITY_GAP' and current[7] == 'DONE'
        assert json.loads(current[8])['physicalRow'] == physical
    else:
        assert current == prior, ('old-event-changed', physical)
events = [r for r in new if r[2] == spec['id']]
assert [r[7] for r in events] == ['PREPARED', 'DONE']
assert len({r[0] for r in new if r[0]}) == sum(bool(r[0]) for r in new)
assert all(not any(r) for r in new[len(old):] if r[2] != spec['id'])
result = {'result': 'PASS', 'oldRowsUnchanged': 1071, 'gapNotes': 49, 'repairEvents': [r[7] for r in events], 'uniqueEventIds': True, 'otherTabsUnchanged': True}
(root / 'session42-repair-assertions.json').write_text(json.dumps(result, indent=2) + '\n', encoding='utf8')
print(json.dumps(result))
