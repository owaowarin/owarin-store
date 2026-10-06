"""Verify the Session42 isolated-sheet result against the post-F05 snapshot."""
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula


root = Path(__file__).parent
before = load_workbook(root / 'attempts/F05/session42-repaired.xlsx', data_only=False)
after = load_workbook(root / 'google-final.xlsx', data_only=False)
values = load_workbook(root / 'google-final.xlsx', data_only=True)


def rows(book, tab):
    return [tuple(('ARRAY_FORMULA', x.text, x.ref) if isinstance(x, ArrayFormula) else x for x in row)
            for row in book[tab].values]


assert before.sheetnames == after.sheetnames
for tab in ('GAME GUIDE BOOKS', 'ADD REQUESTS'):
    assert rows(before, tab) == rows(after, tab), ('unexpected-tab-change', tab)
for tab in ('SALES', 'ORDERS', 'ORDER LINES'):
    old, new = rows(before, tab), rows(after, tab)
    assert new[:len(old)] == old and len(new) == len(old) + 1, ('old-row-change', tab)

old_mag, new_mag = rows(before, 'MAGAZINE'), rows(after, 'MAGAZINE')
assert len(old_mag) == len(new_mag)
changed_mag = [i for i, (a, b) in enumerate(zip(old_mag, new_mag), 1) if a != b]
assert changed_mag == [1252], ('inventory-diff', changed_mag)
mag_header = new_mag[0]
mag_old, mag_new = old_mag[1251], new_mag[1251]
assert mag_new[mag_header.index('Product ID')] == 'W1V34CHECK-1'
assert mag_old[mag_header.index('Status')] == 'Instock' and mag_new[mag_header.index('Status')] == 'Sold'
assert mag_old[mag_header.index('Item UID')] is None and mag_new[mag_header.index('Item UID')]
for i in range(len(mag_header)):
    if i not in (mag_header.index('Status'), mag_header.index('Sold Date'), mag_header.index('Item UID')):
        assert mag_old[i] == mag_new[i], ('inventory-cell-change', i)

old_journal, new_journal = rows(before, 'ORDER REQUESTS'), rows(after, 'ORDER REQUESTS')
for physical, row in enumerate(old_journal, 1):
    if any(x is not None for x in row):
        assert new_journal[physical - 1] == row, ('old-event-change', physical)
new_events = [(i, r) for i, r in enumerate(new_journal, 1)
              if i > 1 and r[0] and (i > len(old_journal) or not old_journal[i - 1][0])]
assert len(new_events) == 28 and [i for i, _ in new_events] == list(range(1124, 1152))
assert len({r[0] for r in new_journal if r[0]}) == sum(bool(r[0]) for r in new_journal)
by_request = {}
for _, row in new_events:
    by_request.setdefault(row[2], []).append(row)
create_id = 'w1-1791153252195-hx9e08l17e6'
confirm_id = 'w1-1791153462270-cw715xj4mec'
assert set(by_request) == {create_id, confirm_id}, ('unexpected-request', set(by_request))
assert [r[7] for r in by_request[create_id]][0:1] == ['PREPARED']
assert [r[7] for r in by_request[confirm_id]][0:1] == ['PREPARED']
assert len(by_request[create_id]) == 11 and len(by_request[confirm_id]) == 17
assert all(by_request[k][-1][7] == 'DONE' for k in by_request)
assert all(r[6] == 'OWA-20261005-01' for rs in by_request.values() for r in rs)

orders = rows(values, 'ORDERS')
lines = rows(values, 'ORDER LINES')
order = [r for r in orders if r[0] == 'OWA-20261005-01']
line = [r for r in lines if r[0] == 'OWA-20261005-01']
assert len(order) == len(line) == 1
assert order[0][1:3] == ('SOLD', 'SHOP') and order[0][7:11] == (390, 50, 10, 440)
assert line[0][2] == mag_new[mag_header.index('Item UID')]
assert line[0][4] == 'W1V34CHECK-1' and line[0][12] == 'SOLD'
assert len([r for r in orders if r[0] == 'OWA-20261004-08' and r[1] == 'PENDING']) == 1
sales = rows(values, 'SALES')
sales_header = sales[0]
sale_lines = [r for r in sales[1:] if r[sales_header.index('Order Line ID')]]
assert len(sale_lines) == 106
assert len({r[sales_header.index('Order Line ID')] for r in sale_lines}) == 106
sale = [r for r in sale_lines if r[sales_header.index('Order ID')] == 'OWA-20261005-01']
assert len(sale) == 1 and sale[0][sales_header.index('Price')] == 390
assert sale[0][sales_header.index('Shipping Cost')] == 10
assert sale[0][sales_header.index('Net Profit')] == 280
assert sale[0][sales_header.index('Order Line ID')] == line[0][1]

result = {'result': 'PASS', 'candidate': json.loads((root / 'revision.json').read_text())['revision'],
          'newOrder': 'OWA-20261005-01', 'createRequest': create_id, 'confirmRequest': confirm_id,
          'newEvents': len(new_events), 'uniqueSalesLines': len(sale_lines),
          'baselineRowsAndUnrelatedTabsPreserved': True, 'originalPending08Preserved': True}
(root / 'session42-final-assertions.json').write_text(json.dumps(result, indent=2) + '\n', encoding='utf8')
print(json.dumps(result))
