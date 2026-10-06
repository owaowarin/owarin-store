import base64, hashlib, json, zlib
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

p = Path(__file__).parent
b = load_workbook(p/'google-before.xlsx', data_only=False)
f = load_workbook(p/'google-final.xlsx', data_only=False)
v = load_workbook(p/'google-final.xlsx', data_only=True)
def norm(x): return (x.text, x.ref) if isinstance(x, ArrayFormula) else x
def rows(w, n): return [[norm(c) for c in r] for r in w[n].values]
def used(rr):
    while rr and not any(x is not None for x in rr[-1]): rr.pop()
    return rr
assert set(f.sheetnames) == set(b.sheetnames) | {'CLIENT'}
counts = {}
for n in b.sheetnames:
    old, new = used(rows(b,n)), used(rows(f,n))
    assert new[:len(old)] == old, ('old cells/formulas changed', n)
    counts[n] = {'before':len(old), 'after':len(new)}
    if n in ['GAME GUIDE BOOKS','W1 DECIMAL REVIEW','ADD REQUESTS']: assert new == old, n
ids = {'OWA-20261005-04','OWA-20261005-05'}
orders = [r for r in rows(v,'ORDERS') if r[0] in ids]
assert len(orders)==2 and all(r[1]=='SOLD' and r[12]==3 and r[6] and r[11] for r in orders)
assert all((r[7],r[8],r[9],r[10])==(390.1,50,10,440.1) for r in orders)
lines = [r for r in rows(v,'ORDER LINES') if r[0] in ids]
assert len(lines)==2 and all(r[12]=='SOLD' for r in lines)
sales = [(i+1,r) for i,r in enumerate(rows(v,'SALES')) if r[0] in ids]
assert len(sales)==2 and len({r[9] for _,r in sales})==2
for row,r in sales:
    assert r[4]==390.1 and r[5]==10
    assert f['SALES'].cell(row,7).value==f'=E{row}-D{row}-F{row}'
client_rows = [r for r in rows(f,'CLIENT')[1:] if any(x is not None for x in r)]
assert len(client_rows)==2 and len({r[6] for r in client_rows})==2
assert {r[2] for r in client_rows}=={'0890000001','0890000003'}
assert {r[4] for r in client_rows}=={'00123','00124'}
assert {r[0] for r in client_rows}=={'=FB INTERNAL SYNTHETIC W2','+FB INTERNAL SYNTHETIC'}
assert {r[5] for r in client_rows}=={'@INTERNAL NEVER PRINT','=2+3 INTERNAL'}
assert all(c.data_type!='f' for row in f['CLIENT'] for c in row)
assert all(isinstance(x,str) for r in client_rows for x in r if x is not None)
mag = [r for r in rows(v,'MAGAZINE') if str(r[1]).startswith('W2V39-')]
assert len(mag)==3 and {r[1]:r[2] for r in mag}=={'W2V39-0':'Sold','W2V39-1':'Sold','W2V39-2':'Instock'}
events = used(rows(f,'ORDER REQUESTS'))
old_events = used(rows(b,'ORDER REQUESTS'))
added = events[len(old_events):]
assert added and all(r[0] and r[2] and r[7] and r[8] for r in added)
latest = {r[2]:r for r in events[1:]}
assert all(r[7]=='DONE' for r in latest.values())
def decode(s):
    if s.startswith('GZIP:'): s=zlib.decompress(base64.b64decode(s[5:]),31).decode()
    return json.loads(s)
requests = []
for k,r in latest.items():
    if any(e[2]==k for e in added):
        snap = decode(r[8])
        requests.append({'id':k,'action':r[4],'state':r[7],'attempt':r[3],'originalPayload':bool(snap.get('payload'))})
all_ids = [r[9] for r in rows(v,'SALES')[1:] if r[9]]
assert len(all_ids)==len(set(all_ids))
assert counts['ORDERS']['after']==counts['ORDERS']['before']+2
assert counts['ORDER LINES']['after']==counts['ORDER LINES']['before']+2
assert counts['SALES']['after']==counts['SALES']['before']+2
new_mag=rows(f,'MAGAZINE')[counts['MAGAZINE']['before']:]
assert sum(any(x is not None for x in r) for r in new_mag)==3
assert [(i+1,r[1]) for i,r in enumerate(rows(v,'MAGAZINE')) if str(r[1]).startswith('W2V39-')]==[(1351,'W2V39-0'),(1352,'W2V39-1'),(1353,'W2V39-2')]
result = {'result':'PASS','oldTabsCellsFormulasAndJournalPrefixPreserved':True,'addedOrders':2,'addedLines':2,'addedSales':2,'clients':2,'newJournalEvents':len(added),'oldJournalPrefixRows':len(old_events),'allLatestRequestsDone':True,'uniqueSalesLineIds':len(all_ids),'counts':counts,'requests':requests,'literalClientValuesAndNoFormulas':True,'xlsxBytes':(p/'google-final.xlsx').stat().st_size,'xlsxSHA256':hashlib.sha256((p/'google-final.xlsx').read_bytes()).hexdigest()}
if (p/'google-native.json').exists():
    native=json.loads((p/'google-native.json').read_text())
    replay=native['replay']
    assert replay['result']=='PASS' and all(r['replayed'] for r in replay['replays'])
    assert native['final']==replay['proof']
    assert len(replay['replays'])==6
    result['sameIdReplays']=6
(p/'google-assertions.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result,indent=2))
