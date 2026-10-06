from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
e=Path(__file__).resolve().parent
before=load_workbook(e/'before/shop-sheet.xlsx',data_only=False)
after=load_workbook(e/'shop-sheet-after.xlsx',data_only=False)
assert after.sheetnames.count('ADD REQUESTS')==1
assert [n for n in after.sheetnames if n!='ADD REQUESTS']==before.sheetnames,after.sheetnames
compared=0
def value(cell):
    v=cell.value if cell is not None else None
    return ('ARRAY_FORMULA',v.text,v.ref) if isinstance(v,ArrayFormula) else v
for name in before.sheetnames:
    a,b=before[name],after[name]
    for row,col in set(a._cells)|set(b._cells):
        assert value(a._cells.get((row,col)))==value(b._cells.get((row,col))),(name,row,col,'existing cell changed')
        compared+=1
j=after['ADD REQUESTS']
headers=['Event ID','Timestamp','Request ID','Attempt','State','Actor','Entry Source','App Version','Sheet','Payload Hash','Row','Product ID','Before','After','Error','Result','Recovery Ref']
assert [j.cell(1,i).value for i in range(1,18)]==headers
assert all(cell.value is None for row in j.iter_rows(min_row=2) for cell in row),'unexpected journal events'
print(f'PASS: {len(before.sheetnames)} original tabs and {compared} values/formulas unchanged; only new ADD REQUESTS with exact17 headers and no events; no inventory/SALES/CLIENT change')
