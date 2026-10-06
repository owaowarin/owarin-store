import json
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).parent
before = load_workbook(root / 'test-sheet-after.xlsx', data_only=False)
after = load_workbook(root / 'test-sheet-after-retry.xlsx', data_only=False)
assert before.sheetnames == after.sheetnames == ['GAME GUIDE BOOKS', 'MAGAZINE', 'ADD REQUESTS']

changes = []
for name in before.sheetnames:
    old, new = before[name], after[name]
    for row in range(1, max(old.max_row, new.max_row) + 1):
        for col in range(1, max(old.max_column, new.max_column) + 1):
            a, b = old.cell(row, col).value, new.cell(row, col).value
            if a == b:
                continue
            assert a is None, (name, row, col, 'existing value changed', a, b)
            assert (name == 'GAME GUIDE BOOKS' and row == 26 or
                    name == 'ADD REQUESTS' and row in (78, 79)), (name, row, col, a, b)
            changes.append((name, row, col))

item = after['GAME GUIDE BOOKS']
assert [item[f'{c}26'].value for c in 'ABCGJKM'] == [
    'Synthetic QA Retry 20261003 A34', 'OWA-GGBS024N00',
    'New Arrival', 'TEST', 301, 101, 296]
for col in 'NOP':
    assert item[f'{col}26'].value.startswith('=IF('), col
assert item['V26'].value.startswith('=IFERROR(__xludf.DUMMYFUNCTION(')
journal = after['ADD REQUESTS']
request = 'add-1790981261601-tr7vx25ksx'
assert [journal[f'E{r}'].value for r in (78, 79)] == ['PREPARED', 'DONE']
assert all(journal[f'C{r}'].value == request for r in (78, 79))
assert all(journal[f'K{r}'].value == 26 and journal[f'L{r}'].value == 'OWA-GGBS024N00' for r in (78, 79))
result = json.loads(journal['P79'].value)
assert result['itemSnapshot']['name'] == item['A26'].value
assert result['itemSnapshot']['productId'] == item['B26'].value
assert result['itemSnapshot']['cost'] == 101 and result['itemSnapshot']['price'] == 296
assert result['success'] is True
assert journal['A80'].value is None and item['A27'].value is None
print(f'PASS: retry preserved original ID; only GGB row 26 and journal rows 78-79 changed ({len(changes)} populated cells); no duplicate or prior-cell change')
