import json
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).parent
before = load_workbook(root / 'test-sheet-before.xlsx', data_only=False)
after = load_workbook(root / 'test-sheet-after.xlsx', data_only=False)
assert before.sheetnames == after.sheetnames == ['GAME GUIDE BOOKS', 'MAGAZINE', 'ADD REQUESTS']

changes = []
for name in before.sheetnames:
    old, new = before[name], after[name]
    for row in range(1, max(old.max_row, new.max_row) + 1):
        for col in range(1, max(old.max_column, new.max_column) + 1):
            a, b = old.cell(row, col).value, new.cell(row, col).value
            if a == b:
                continue
            assert a is None, (name, row, col, 'existing value changed', a, b)
            assert (name == 'GAME GUIDE BOOKS' and row == 25 or
                    name == 'ADD REQUESTS' and row in (76, 77)), (name, row, col, a, b)
            changes.append((name, row, col))

item = after['GAME GUIDE BOOKS']
assert [item[f'{c}25'].value for c in 'ABCGJKM'] == [
    'Synthetic QA Integrity 20261003 A33', 'OWA-GGBS023N00',
    'New Arrival', 'TEST', 300, 99, 295]
for col in 'NOP':
    assert item[f'{col}25'].value.startswith('=IF('), col
assert item['V25'].value.startswith('=IFERROR(__xludf.DUMMYFUNCTION(')
journal = after['ADD REQUESTS']
request = 'add-1790980979593-ggaya8nf2s7'
assert [journal[f'E{r}'].value for r in (76, 77)] == ['PREPARED', 'DONE']
assert all(journal[f'C{r}'].value == request for r in (76, 77))
assert all(journal[f'K{r}'].value == 25 and journal[f'L{r}'].value == 'OWA-GGBS023N00' for r in (76, 77))
result = json.loads(journal['P77'].value)
assert result['itemSnapshot']['name'] == item['A25'].value
assert result['itemSnapshot']['productId'] == item['B25'].value
assert result['itemSnapshot']['cost'] == 99 and result['itemSnapshot']['price'] == 295
assert result['success'] is True
assert any(name == 'GAME GUIDE BOOKS' for name, _, _ in changes)
assert any(name == 'ADD REQUESTS' for name, _, _ in changes)
print(f'PASS: only GGB row 25 and journal rows 76-77 changed ({len(changes)} populated cells); Cost/Price/formulas/snapshot read back; MAG and prior cells unchanged')
