import json,hashlib,difflib
from pathlib import Path
from openpyxl import load_workbook
p=Path(__file__).parent;root=p.parents[2]
sources=json.loads((p/'sources-before.json').read_text(encoding='utf-8-sig'))
correct=json.loads((p/'manifest-before.json').read_text(encoding='utf-8-sig'))['source']
assert json.loads(correct)['runtimeVersion']=='V8'
stale=p/'before/source/appsscript.json'
if stale.exists() and stale.read_text(encoding='utf-8')!=correct.replace('\r\n','\n'):
 attempt=p/'before/manifest-stale-attempt.txt'
 assert not attempt.exists()
 stale.rename(attempt)
sources['appsscript.json']=correct
assert set(sources)=={'Code.gs','webapp.gs','Index.html','FbAlbum.gs','R2Upload.gs','P1Journal.gs','appsscript.json'}
dest=p/'before/source';dest.mkdir(parents=True,exist_ok=True)
sha=lambda s:hashlib.sha256(s.replace('\r\n','\n').encode()).hexdigest().upper()
for n,s in sources.items():
 f=dest/n
 if f.exists():assert f.read_text(encoding='utf-8')==s.replace('\r\n','\n')
 else:f.write_text(s,encoding='utf-8',newline='')
manifest=json.loads(sources['appsscript.json']);assert manifest['runtimeVersion']=='V8'
w=load_workbook(p/'shop-before.xlsx',data_only=False)
schema={sh.title:{'rows':sh.max_row,'columns':sh.max_column,'headers':[c.value for c in sh[1]]} for sh in w}
(p/'shop-schema-before.json').write_text(json.dumps(schema,ensure_ascii=False,default=str,indent=2)+'\n',encoding='utf-8')
mapping={'Code.gs':'Code_v31.gs','webapp.gs':'WebApp_v31.gs','Index.html':'Index.html','P1Journal.gs':'P1Journal.gs'}
diffs={}
for n,old in mapping.items():
 prior=(root/'03 Apps Script/Web App'/old).read_text(encoding='utf-8').replace('\r\n','\n');current=sources[n].replace('\r\n','\n')
 d=''.join(difflib.unified_diff(prior.splitlines(True),current.splitlines(True),fromfile='local-v31/'+n,tofile='fresh-production/'+n))
 (p/(n+'.before.diff')).write_text(d,encoding='utf-8');diffs[n]=len(d.splitlines())
out={'result':'PASS','sources':{n:{'chars':len(s),'LFsha256':sha(s)} for n,s in sources.items()},'workbookBytes':(p/'shop-before.xlsx').stat().st_size,'workbookSha256':hashlib.sha256((p/'shop-before.xlsx').read_bytes()).hexdigest(),'tabs':list(schema),'changedFromLocalV31DiffLines':diffs,'manifestRuntime':manifest['runtimeVersion']}
(p/'backup-verification.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'result':'PASS','sourceFiles':len(sources),'tabs':len(schema),'bytes':out['workbookBytes'],'diffLines':diffs,'CLIENTheaders':schema['CLIENT']['headers'],'GGBcolumns':schema['GAME GUIDE BOOKS']['columns'],'MAGcolumns':schema['MAGAZINE']['columns']}))
