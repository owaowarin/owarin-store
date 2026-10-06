import csv,json,hashlib,shutil
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
p=Path(__file__).parent
def log(action,before,after,detail):
 with (p/'changes.csv').open('a',encoding='utf-8',newline='') as f: csv.writer(f).writerow([datetime.now().astimezone().isoformat(),'W3-PROD-20261006-01',action,before,after,detail])
src=Path('C:/Users/JIN/AppData/Local/Temp/browser-use/exports/OWARIN STORE — BACKUP before W1-W2 — 2026-10-06-2c40ce48-03d9-45c1-bf56-d7e43c6e9e2b.xlsx')
target=p/'shop-backup-copy.xlsx'
assert src.is_file() and not target.exists()
log('COPY-DRYRUN',str(src),'ABSENT','Native backup export for stored-cell comparison')
shutil.copy2(src,target)
log('COPY',str(src),hashlib.sha256(target.read_bytes()).hexdigest(),'Native-copy export saved')
a=load_workbook(p/'shop-before.xlsx',data_only=False);b=load_workbook(target,data_only=False)
assert a.sheetnames==b.sheetnames
differences=[]
for name in a.sheetnames:
 sa,sb=a[name],b[name]
 for row in range(1,max(sa.max_row,sb.max_row)+1):
  for col in range(1,max(sa.max_column,sb.max_column)+1):
   ca,cb=sa.cell(row,col),sb.cell(row,col)
   if ca.value!=cb.value or ca.data_type!=cb.data_type:differences.append({'sheet':name,'cell':ca.coordinate,'beforeType':ca.data_type,'afterType':cb.data_type})
out={'result':'PASS' if not differences else 'REVIEW','tabs':len(a.sheetnames),'storedCellDifferences':differences,'copySha256':hashlib.sha256(target.read_bytes()).hexdigest(),'copyId':'1OegftlJuM0glHqymOcMU0G3KhiF_AoYlNR3m1tDoZgY','access':'Restricted; only owner; UI readback backup-copy.json','note':'Native copy preserves bound script/comments; workbook comparison covers stored values/formulas, not recalculated external caches or trigger execution.'}
(p/'copy-verification.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
log('VERIFY','Native backup copy',out['result'],f'{len(a.sheetnames)} tabs; {len(differences)} stored-cell differences')
print(json.dumps({'result':out['result'],'tabs':out['tabs'],'differences':len(differences)}))
