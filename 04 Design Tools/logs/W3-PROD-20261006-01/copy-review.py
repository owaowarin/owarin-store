import json,collections
from pathlib import Path
from openpyxl import load_workbook
p=Path(__file__).parent
a=load_workbook(p/'shop-before.xlsx',data_only=False);b=load_workbook(p/'shop-backup-copy.xlsx',data_only=False)
def val(c):
 v=c.value
 return (c.data_type,v.text if hasattr(v,'text') else v)
out={}
for n in a.sheetnames:
 sa,sb=a[n],b[n];changed=[];types=collections.Counter()
 for key in set(sa._cells)|set(sb._cells):
  ca,cb=sa.cell(*key),sb.cell(*key)
  if val(ca)!=val(cb):changed.append(ca.coordinate);types[str((ca.data_type,cb.data_type))]+=1
 if changed:out[n]={'cells':len(changed),'types':dict(types),'examples':changed[:8]}
(p/'copy-review.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(out,ensure_ascii=False))
