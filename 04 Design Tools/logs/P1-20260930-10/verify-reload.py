from pathlib import Path
from openpyxl import load_workbook

here = Path(__file__).resolve().parent
before = load_workbook(here / 'test-sheet-before.xlsx', data_only=False)
after = load_workbook(here / 'test-sheet-after.xlsx', data_only=False)
assert before.sheetnames == after.sheetnames == ['GAME GUIDE BOOKS', 'MAGAZINE', 'ADD REQUESTS']

differences = {}
for name in before.sheetnames:
    old, new = before[name], after[name]
    changed = []
    for row in range(1, max(old.max_row, new.max_row) + 1):
        for col in range(1, max(old.max_column, new.max_column) + 1):
            a, b = old.cell(row, col).value, new.cell(row, col).value
            if a != b:
                changed.append((row, col, a, b))
    differences[name] = changed
    assert all(a is None for _, _, a, _ in changed), f'{name}: existing value/formula changed'

ggb = differences['GAME GUIDE BOOKS']
journal = differences['ADD REQUESTS']
assert ggb and {row for row, *_ in ggb} == {24}
assert not differences['MAGAZINE']
assert journal and {row for row, *_ in journal} == {74, 75}
g = after['GAME GUIDE BOOKS']
j = after['ADD REQUESTS']
request_id = 'add-1790721738709-k0dwjsg2psd'
assert (g['A24'].value, g['B24'].value, g['C24'].value) == (
    'Synthetic QA Reload 20260930 A32', 'OWA-GGBS022N00', 'New Arrival')
assert g['A25'].value is None
assert [j[f'C{row}'].value for row in (74, 75)] == [request_id, request_id]
assert [j[f'E{row}'].value for row in (74, 75)] == ['PREPARED', 'DONE']
assert [j[f'D{row}'].value for row in (74, 75)] == [1, 1]
assert [j[f'K{row}'].value for row in (74, 75)] == [24, 24]
assert [j[f'L{row}'].value for row in (74, 75)] == ['OWA-GGBS022N00'] * 2
assert j['C76'].value is None
print(f'PASS: one GGB row24, PREPARED/DONE rows74-75, prior values/formulas unchanged; changes GGB={len(ggb)} MAG=0 journal={len(journal)}')
