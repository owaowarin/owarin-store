import json
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).parent
before = load_workbook(root / 'test-sheet-before.xlsx', data_only=False)
assert before.sheetnames == ['GAME GUIDE BOOKS', 'MAGAZINE', 'ADD REQUESTS']
assert before['GAME GUIDE BOOKS']['A26'].value == 'Synthetic QA Retry 20261003 A34'
assert before['GAME GUIDE BOOKS']['A27'].value is None
assert before['GAME GUIDE BOOKS']['K1'].value == 'Cost'
assert before['ADD REQUESTS']['C79'].value == 'add-1790981261601-tr7vx25ksx'
assert before['ADD REQUESTS']['A80'].value is None
if not (root / 'test-sheet-after.xlsx').exists():
    print('PASS: fresh before export matches guarded GGB row26/journal row79 and Cost header')
    raise SystemExit(0)
after = load_workbook(root / 'test-sheet-after.xlsx', data_only=False)
assert before.sheetnames == after.sheetnames
changed = []
for name in before.sheetnames:
    old, new = before[name], after[name]
    for row in range(1, max(old.max_row, new.max_row) + 1):
        for col in range(1, max(old.max_column, new.max_column) + 1):
            a, b = old.cell(row, col).value, new.cell(row, col).value
            if a == b:
                continue
            assert a is None, (name, row, col, 'previous cell changed', a, b)
            assert (name == 'GAME GUIDE BOOKS' and row in (27,28,29) or
                    name == 'ADD REQUESTS' and 80 <= row <= 87), (name,row,col,a,b)
            changed.append((name,row,col))
items, journal = after['GAME GUIDE BOOKS'], after['ADD REQUESTS']
tag = 'p1-review-20261003-02-'
for row, label, seq in [(27,'R2','025'),(28,'R1','026'),(29,'R4','027')]:
    assert items[f'A{row}'].value == 'Synthetic QA Astra 20261003 ' + label
    assert items[f'B{row}'].value == 'OWA-GGBS' + seq + 'N00'
    assert items[f'K{row}'].value == 99 and items[f'M{row}'].value == 299
    assert items[f'N{row}'].value.startswith('=IF(')
    assert items[f'P{row}'].value.startswith('=IF(')
    assert items[f'V{row}'].value.startswith('=IFERROR(__xludf.DUMMYFUNCTION(')
    if label == 'R4':
        assert items[f'O{row}'].value is None
    else:
        assert items[f'O{row}'].value.startswith('=IF(')
expected = [(80,'R2','PREPARED'),(81,'R2','DONE'),(82,'R1','PREPARED'),(83,'R1','DONE'),
            (84,'R1','DONE'),(85,'R4','PREPARED'),(86,'R4','ERROR'),(87,'R4','ERROR')]
for row, label, state in expected:
    assert journal[f'C{row}'].value == tag + label
    assert journal[f'E{row}'].value == state
for row in (81,83,84):
    result=json.loads(journal[f'P{row}'].value)
    assert result['itemSnapshot']['cost'] == 99 and result['itemSnapshot']['price'] == 299
assert 'formula marketplace' in journal['O86'].value
assert 'Retry blocked pending recovery' in journal['O87'].value
assert items['A30'].value is None and journal['A88'].value is None
print(f'PASS: {len(changed)} added cells only; GGB rows27-29/journal80-87; R2 retry DONE; R1 replay DONE; R4 ERROR/no duplicate; all prior cells and MAG unchanged')
