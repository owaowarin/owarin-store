"""Read-only extraction: no CLIENT data values, no workbook save."""
import collections
import json
import os
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parent.parent
source = Path(os.environ['TEMP']) / 'OWARIN-P0-20260928-01-latest.xlsx'
out = {}
book = openpyxl.load_workbook(source, read_only=True, data_only=False)
for name in ['GAME GUIDE BOOKS', 'MAGAZINE']:
    sheet = book[name]
    counts = collections.Counter()
    statuses, missing, formulas = set(), [], []
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    for row in sheet.iter_rows(min_row=3, max_col=3):
        title, pid, status = [cell.value for cell in row]
        if not title:
            continue
        if status:
            statuses.add(str(status))
        if not pid:
            missing.append(row[1].coordinate)
        elif row[1].data_type == 'f':
            formulas.append(row[1].coordinate)
        else:
            counts[str(pid).strip()] += 1
    out[name] = dict(statusValues=sorted(statuses), missingPIDCells=missing,
                     formulaPIDCells=formulas, duplicatePIDGroups=sum(n > 1 for n in counts.values()),
                     immutableUIDHeaderPresent='Item UID' in headers)
sheet = book['SALES']
g2 = sheet['G2'].value
out['SALES'] = dict(G2=g2 if isinstance(g2, str) else dict(text=g2.text, ref=g2.ref),
    formulaCells=[cell.coordinate for row in sheet.iter_rows(min_col=7, max_col=7) for cell in row if cell.data_type == 'f'],
    lastNonemptyOrderIDRow=max([cell.row for row in sheet.iter_rows(min_col=1, max_col=1) for cell in row if cell.value] + [0]))
book.close()
result = json.dumps(out, ensure_ascii=False, indent=2)
(root / 'identity-formula-summary.json').write_text(result, encoding='utf-8')
print(result)
