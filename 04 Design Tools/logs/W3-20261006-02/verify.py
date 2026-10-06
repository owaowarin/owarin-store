import base64,hashlib,json,zlib
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
p=Path(__file__).parent
def cells(f,data=False):
 w=load_workbook(f,data_only=data);out={}
 for sh in w:
  rows=[[{'formula':v.text,'ref':v.ref} if isinstance(v,ArrayFormula) else v.isoformat() if isinstance(v,datetime) else v for v in r] for r in sh.values]
  while rows and not any(v is not None for v in rows[-1]):rows.pop()
  for r in rows:
   while r and r[-1] is None:r.pop()
  out[sh.title]=rows
 return out
def decode(s):return json.loads(zlib.decompress(base64.b64decode(s[5:]),31).decode() if s.startswith('GZIP:') else s)
def requests(rows):return {r[2]:{'id':r[2],'action':r[4],'state':r[7],'attempt':r[3],'snap':decode(r[8])} for r in rows[1:] if r}
old,ui=cells(p/'google-before.xlsx'),cells(p/'google-after-ui.xlsx')
assert set(old)==set(ui)
for n in set(old)-{'MAGAZINE','SALES','CLIENT','ORDERS','ORDER LINES','ORDER REQUESTS'}:assert old[n]==ui[n],n
for n,added in [('SALES',2),('ORDERS',2),('ORDER LINES',2)]:
 assert ui[n][:len(old[n])]==old[n],n+' prefix'
 assert len(ui[n])-len(old[n])==added,n+' count'
assert ui['CLIENT'][:len(old['CLIENT'])]==old['CLIENT']
assert sum(bool(r) for r in ui['CLIENT'])-sum(bool(r) for r in old['CLIENT'])==1
assert all(not r for r in ui['CLIENT'][len(old['CLIENT']):-1]),'Only blank capacity rows may precede new client'
assert ui['ORDER REQUESTS'][:len(old['ORDER REQUESTS'])]==old['ORDER REQUESTS']
magdiff=[]
for i,(a,b) in enumerate(zip(old['MAGAZINE'],ui['MAGAZINE'])):
 if a==b:continue
 assert a[1] in ['W2V39-2','W1V33CHECK-0']
 for col in range(max(len(a),len(b))):
  av=a[col] if col<len(a) else None;bv=b[col] if col<len(b) else None
  if av!=bv:assert col in [2,10,17,23],(i+1,col+1,av,bv)
 assert b[2]=='Sold';magdiff.append({'row':i+1,'sku':a[1]})
assert len(magdiff)==2 and len(old['MAGAZINE'])==len(ui['MAGAZINE'])
orders=ui['ORDERS'][-2:];assert [r[0] for r in orders]==['OWA-20261006-01','OWA-20261006-02']
assert orders[0][1:3]==['SOLD','SHOP'] and orders[0][7:11]==[390.1,50,10,440.1]
assert orders[1][1:3]==['SOLD','SHOPEE'] and orders[1][7:11]==[500,0,0,500]
for i,r in enumerate(orders):
 label=json.loads(r[11]);assert label['postalCode']==['00201','00203'][i];assert label['phone']==['0890000201','0890000203'][i];assert r[12]==2
lines=ui['ORDER LINES'][-2:];assert [r[8:10] for r in lines]==[[390.1,390.1],[500,300]] and all(r[12]=='SOLD' for r in lines)
assert len({r[1] for r in lines})==2
sales=cells(p/'google-after-ui.xlsx',True)['SALES'][-2:]
assert [r[0] for r in sales]==[r[0] for r in orders] and [r[9] for r in sales]==[r[1] for r in lines]
assert [r[3:7] for r in sales]==[[100,390.1,10,280.1],[100,300,0,200]],'Ledger/cost/subsidy/profit'
client=ui['CLIENT'][-1];assert client[:6]==['=FB W3 SYNTHETIC','ผู้ซื้อ W3 SHOPEE จำลอง','0890000202','77 ถนนลูกค้าทดสอบ กรุงเทพ','00202','@PRIVATE W3 NEVER PRINT']
formula_client=load_workbook(p/'google-after-ui.xlsx',data_only=False)['CLIENT'];assert all(formula_client.cell(len(ui['CLIENT']),c).data_type!='f' for c in range(1,9))
req=requests(ui['ORDER REQUESTS']);newreq=[r for r in req.values() if r['snap'].get('orderId') in [o[0] for o in orders]]
assert len(newreq)==5 and all(r['state']=='DONE' for r in newreq)
out={'result':'PASS','ui':{'channels':['SHOP','SHOPEE'],'orders':[r[0] for r in orders],'newSales':2,'newClient':1,'oldPrefixesPreserved':True,'changedInventory':magdiff,'requests':[{'id':r['id'],'action':r['action'],'state':r['state']} for r in newreq],'journalAddedEvents':len(ui['ORDER REQUESTS'])-len(old['ORDER REQUESTS'])},'print':'DEFERRED — ยังไม่ได้ทดสอบ; NOT PASS'}
if (p/'google-final.xlsx').exists():
 before,final,after=cells(p/'google-before-migration.xlsx'),cells(p/'google-final.xlsx'),cells(p/'google-after-replay.xlsx')
 assert set(before)==set(final)==set(ui)|{'W3 LEGACY CLIENT','W3 MIGRATION REQUESTS'}
 for n in ui:
  if n!='ORDER REQUESTS':assert before[n]==ui[n],n+' fixture initialization'
  assert before[n]==final[n],n+' migration changed original'
 assert before['ORDER REQUESTS'][:len(ui['ORDER REQUESTS'])]==ui['ORDER REQUESTS']
 init=requests(before['ORDER REQUESTS']);assert init['qa-w3-native-fixture-20261006']['state']=='DONE'
 assert len(before['ORDER REQUESTS'])-len(ui['ORDER REQUESTS'])==5
 legacybefore,legacyfinal=before['W3 LEGACY CLIENT'],final['W3 LEGACY CLIENT']
 assert len(legacybefore)==len(legacyfinal)==5 and len(legacybefore[0])==6 and len(legacyfinal[0])==8
 for a,b in zip(legacybefore,legacyfinal):assert a==b[:6],(a,b)
 nativebefore=cells(p/'google-before-migration.xlsx',True)['W3 LEGACY CLIENT'];nativefinal=cells(p/'google-final.xlsx',True)['W3 LEGACY CLIENT']
 assert nativebefore[1][5]==2 and legacybefore[1][5]=='=1+1';assert nativefinal[1][5]==2
 assert legacybefore[1][2]=='0890000101' and legacybefore[1][4]=='00123' and legacybefore[4][2]==890000003 and legacybefore[4][4]==123
 ids=[(r+[None]*8)[6:8] for r in legacyfinal[1:]];assert ids[1]==[None,None];assert len({r[0] for r in ids if r[0]})==3
 failed=json.loads((p/'native-failure.json').read_text());retry=json.loads((p/'native-retry.json').read_text());replay=json.loads((p/'native-replay.json').read_text())
 assert failed['result']=='EXPECTED_FAILURE' and failed['state']=='NEEDS_REVIEW' and 'RECOVERY_REQUIRED' in failed['error']
 assert retry['result']=='PASS' and retry['state']=='DONE' and retry['attempt']==2
 assert retry['allocatedIds']==failed['allocatedIds'] and retry['allocatedRevision']==failed['allocatedRevision']
 assert ids==[[id or None,retry['allocatedRevision'] if id else None] for id in retry['allocatedIds']]
 assert replay['result']=='PASS' and replay['migration']['replayed'] and replay['ui']['result']=='PASS' and len(replay['ui']['replays'])==5
 assert {r['id'] for r in replay['ui']['replays']}=={r['id'] for r in newreq}
 assert final==after,'DONE replay changed native values/formulas'
 alias=requests(final['W3 MIGRATION REQUESTS']);assert len(alias)==1 and alias['qa-w2-client-schema-20261005']['state']=='DONE'
 out['nativeMigration']={'result':'PASS','originalAFValuesFormulas':True,'originalNineTabsUnchanged':True,'sameIdsRevisionRetry':True,'blankRowDistinctIds':True,'allocatedIds':retry['allocatedIds'],'allocatedRevision':retry['allocatedRevision'],'journalEvents':len(final['W3 MIGRATION REQUESTS'])-1}
 out['nativeReplay']={'result':'PASS','migrationReplay':True,'uiOriginalRequests':5,'allSheetValuesFormulasUnchanged':True}
 out['mainJournalRows']=len(final['ORDER REQUESTS'])-1
 out['exports']={n:hashlib.sha256((p/n).read_bytes()).hexdigest() for n in ['google-before.xlsx','google-after-ui.xlsx','google-before-migration.xlsx','google-final.xlsx','google-after-replay.xlsx']}
(p/'verification.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps(out,ensure_ascii=False))
