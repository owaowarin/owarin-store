from openpyxl import load_workbook

def rows(path, sheet):
    book = load_workbook(path, read_only=True, data_only=False)
    return list(book[sheet].iter_rows(values_only=True))

before = '04 Design Tools/logs/P1-20260930-03/test-sheet-before.xlsx'
ggb = rows(before, 'GAME GUIDE BOOKS')
mag = rows(before, 'MAGAZINE')
journal = rows(before, 'ADD REQUESTS')
print('before GGB named', sum(bool(row[0]) for row in ggb[2:]), 'MAG named', sum(bool(row[0]) for row in mag[2:]))
print('before journal events', sum(bool(row[0]) for row in journal[1:]))
print('GGB row22', ggb[21][:3], 'journal row70', journal[69][:5] if len(journal) > 69 else None)
