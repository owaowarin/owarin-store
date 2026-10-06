import base64, hashlib, json, math, zlib
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
p=Path(__file__).parent
b=load_workbook(p/'google-before-v38.xlsx',data_only=False)
f=load_workbook(p/'google-final-v38.xlsx',data_only=False)
v=load_workbook(p/'google-final-v38.xlsx',data_only=True)
assert set(b.sheetnames)==set(f.sheetnames)
def norm(x): return (x.text,x.ref) if isinstance(x,ArrayFormula) else x
def rows(w,n): return [[norm(c) for c in r] for r in w[n].values]
def used(rr):
 while rr and not any(x is not None for x in rr[-1]): rr.pop()
 return rr
for n in b.sheetnames:
 old=rows(b,n);new=rows(f,n)
 if n=='ORDER REQUESTS':old=used(old)
 assert new[:len(old)]==old,('old cells changed',n)
 if n in ['GAME GUIDE BOOKS','W1 DECIMAL REVIEW','ADD REQUESTS']:assert used(new)==used(old),n
native=json.loads((p/'google-native-decimal.json').read_text())
replay=json.loads((p/'google-native-replay.json').read_text())
assert len(native['orders'])==2 and len(native['ledger'])==3
assert len(native['requests'])==4 and all(r['state']=='DONE' for r in native['requests'])
assert len(replay)==4 and all(r['replayed'] for r in replay)
order_ids={r[0] for r in native['orders']}
orders=[r for r in rows(v,'ORDERS') if r[0] in order_ids]
assert len(orders)==2 and all(r[1]=='SOLD' and r[2]=='SHOP' and r[12]==2 for r in orders)
assert sorted((r[7],r[8],r[9],r[10]) for r in orders)==[(0.3,60,10.02,60.3),(390.1,50,10.02,440.1)]
lines=[r for r in rows(v,'ORDER LINES') if r[0] in order_ids]
assert len(lines)==3 and all(r[12]=='SOLD' for r in lines)
assert {r[4] for r in lines}=={'W1V38DEC-0','W1V38DEC-1','W1V38DEC-2'}
sales=[(i+1,r) for i,r in enumerate(rows(v,'SALES')) if r[0] in order_ids]
assert len(sales)==3 and len({r[9] for _,r in sales})==3
expected={'W1V38DEC-0':(390.1,100.05,10.02,280.03),'W1V38DEC-1':(0.1,0.05,10.02,-9.97),'W1V38DEC-2':(0.2,0.05,0,0.15)}
for row,r in sales:
 price,cost,subsidy,profit=expected[r[8]]
 assert r[4]==price and r[3]==cost and r[5]==subsidy and math.isclose(r[6],profit,abs_tol=1e-7,rel_tol=0)
 assert f['SALES'].cell(row,7).value==f'=E{row}-D{row}-F{row}'
 assert r[7]==r[8]
 assert sum(l['lineId']==r[9] and l['row']==row for l in native['ledger'])==1
mag=[r for r in rows(v,'MAGAZINE') if str(r[1]).startswith('W1V38DEC-')]
assert len(mag)==3 and all(r[2]=='Sold' and r[17] and r[23] for r in mag)
assert {r[1]:r[8] for r in mag}=={'W1V38DEC-0':100.05,'W1V38DEC-1':0.05,'W1V38DEC-2':0.05}
old_events=used(rows(b,'ORDER REQUESTS'));events=used(rows(f,'ORDER REQUESTS'))
added=events[len(old_events):];assert added and all(r[0] and r[2] and r[7] and r[8] for r in added)
latest={r[2]:r for r in events[1:]};assert all(r[7]=='DONE' for r in latest.values())
assert {r[2] for r in added}=={'qa-v38-decimal-fixture-20261005'}|{r['requestId'] for r in native['requests']}
assert all(r[13]=='Code v38 / WebApp v38' for r in added)
all_line_ids=[r[9] for r in rows(v,'SALES')[1:] if r[9]];assert len(all_line_ids)==len(set(all_line_ids))
assert len(used(rows(f,'ORDERS')))==len(used(rows(b,'ORDERS')))+2
assert len(used(rows(f,'ORDER LINES')))==len(used(rows(b,'ORDER LINES')))+3
assert len([r for r in rows(v,'SALES') if r[9]])==len([r for r in rows(b,'SALES') if r[9]])+3
result={'revision':json.loads((p/'revision.json').read_text())['revision'],'oldTabsCellsAndEventsPreserved':True,'addedOrders':2,'addedLines':3,'addedSales':3,'newJournalEvents':len(added),'oldJournalPrefixRows':len(old_events),'allLatestRequestsDone':True,'uniqueSalesLineIds':len(all_line_ids),'nativeScratch':native['nativeScratch'],'orders':[{'orderId':r[0],'subtotal':r[7],'customerShipping':r[8],'subsidy':r[9],'total':r[10],'status':r[1]} for r in orders],'nativeLedger':native['ledger'],'sameIdReplays':replay,'finalXlsxBytes':(p/'google-final-v38.xlsx').stat().st_size,'finalXlsxSHA256':hashlib.sha256((p/'google-final-v38.xlsx').read_bytes()).hexdigest()}
(p/'google-v38-assertions.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result,indent=2))
