import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
p=Path(__file__).parent
b=load_workbook(p/'google-before.xlsx',data_only=False)
f=load_workbook(p/'google-decimal-final.xlsx',data_only=False)
v=load_workbook(p/'google-decimal-final.xlsx',data_only=True)
def rows(w,n):
 return [[(x.text,x.ref) if isinstance(x,ArrayFormula) else x for x in row] for row in w[n].values]
assert set(f.sheetnames)==set(b.sheetnames)|{'W1 DECIMAL REVIEW'}
for n in b.sheetnames: assert rows(b,n)==rows(f,n),n
assert f['W1 DECIMAL REVIEW']['D2'].value=='=A2-B2-C2'
value=v['W1 DECIMAL REVIEW']['D2'].value
assert f['W1 DECIMAL REVIEW']['C5'].value=='=A5+B5'
assert v['W1 DECIMAL REVIEW']['C5'].value==0.3
assert v['W1 DECIMAL REVIEW']['D5'].value==0.3
result={'formula':'=A2-B2-C2','inputs':[v['W1 DECIMAL REVIEW'][a].value for a in ('A2','B2','C2')],'googleExportCachedValue':value,'jsExpression':390.1-100.05-10.02,'sameDouble':value==390.1-100.05-10.02,'googleSubtotal':v['W1 DECIMAL REVIEW']['C5'].value,'directStoredBinarySum':v['W1 DECIMAL REVIEW']['D5'].value,'existingTabsUnchanged':True}
(p/'google-decimal-final-result.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result))
