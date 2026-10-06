from openpyxl import load_workbook

base = '04 Design Tools/logs/P1-20260930-04/'
before = load_workbook(base + 'test-sheet-before.xlsx', read_only=True, data_only=False)
after = load_workbook(base + 'test-sheet-after.xlsx', read_only=True, data_only=False)

def stable(value):
    if value.__class__.__name__ == 'ArrayFormula':
        return ('ArrayFormula', value.text, value.ref)
    return value

def rows(book, name):
    return [[stable(value) for value in row] for row in book[name].iter_rows(values_only=True)]

assert before.sheetnames == after.sheetnames
old_ggb, new_ggb = rows(before, 'GAME GUIDE BOOKS'), rows(after, 'GAME GUIDE BOOKS')
old_mag, new_mag = rows(before, 'MAGAZINE'), rows(after, 'MAGAZINE')
old_j, new_j = rows(before, 'ADD REQUESTS'), rows(after, 'ADD REQUESTS')
assert old_ggb[:21] == new_ggb[:21], 'existing GGB values/formulas changed'
assert old_mag == new_mag, 'MAGAZINE values/formulas changed'
assert old_j[:69] == new_j[:69], 'existing journal events changed'
assert not old_ggb[21][0] and new_ggb[21][0] == 'Synthetic QA SavedSource 20260930 A30'
assert new_ggb[21][1:3] == ['OWA-GGBS020N00', 'New Arrival']
assert new_ggb[21][6] == 'TEST' and new_ggb[21][9:11] == [300, 100] and new_ggb[21][12] == 292
assert not new_ggb[22][0], 'next GGB row is not empty'
new_events = new_j[69:71]
assert [row[4] for row in new_events] == ['PREPARED', 'DONE']
assert new_events[0][2] == new_events[1][2]
assert [row[3] for row in new_events] == [1, 1]
assert [row[10:12] for row in new_events] == [[22, 'OWA-GGBS020N00']] * 2
assert not new_j[71][0] if len(new_j) > 71 else True
print('PASS existing GGB rows 1:21, all MAGAZINE, journal rows 1:69 unchanged')
print('GGB row22', new_ggb[21][:13])
print('journal rows70:71', [(row[2], row[3], row[4], row[10], row[11], row[14]) for row in new_events])
print('GGB named', sum(bool(row[0]) for row in new_ggb[2:]), 'MAG named', sum(bool(row[0]) for row in new_mag[2:]))
print('journal events', sum(bool(row[0]) for row in new_j[1:]))
