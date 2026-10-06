from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
import json
p=Path(__file__).resolve().parent
w=load_workbook(p/'shop-schema-before.xlsx',data_only=False)
out={}
for n in ['GAME GUIDE BOOKS','MAGAZINE','SALES','CLIENT','ADD REQUESTS']:
 s=w[n]
 out[n]={'headers':[c.value for c in s[1]],'formulas_row2':[{ 'col':c.column,'formula':c.value.text if isinstance(c.value,ArrayFormula) else c.value,'ref':c.value.ref if isinstance(c.value,ArrayFormula) else None} for c in s[2] if c.data_type=='f' or isinstance(c.value,ArrayFormula)],'tables':[{ 'name':t.name,'ref':t.ref,'totals':t.totalsRowCount} for t in s.tables.values()]}
out['new_tabs_absent']=all(n not in w.sheetnames for n in ['ORDERS','ORDER LINES','ORDER REQUESTS'])
(p/'schema.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(out,ensure_ascii=False,indent=2))
